from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from langchain.vectorstores import FAISS
from langchain.embeddings import OpenAIEmbeddings
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.docstore.document import Document
from langchain_openai import ChatOpenAI
from langchain_anthropic import ChatAnthropic
from typing import List
from pydantic import BaseModel
import os
import logging
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from dotenv import load_dotenv

# Configurazione logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

# Caricamento chiavi e configurazione iniziale
load_dotenv()
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY")

if not OPENAI_API_KEY:
    logger.error("La chiave API di OpenAI non è configurata correttamente nel file .env")
    raise RuntimeError("Configurazione API non valida")

# Scegli il modello attivo
ACTIVE_MODEL = "gpt-4o-mini"  # Modello attivo: gpt-4o-mini di OpenAI

# Configurazione embeddings e database
embeddings = OpenAIEmbeddings(openai_api_key=OPENAI_API_KEY)
UPLOAD_FOLDER = "./uploads"
DB_FOLDER = "./db"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(DB_FOLDER, exist_ok=True)

try:
    db_path = "./db/my_database"
    faiss_index = FAISS.load_local(db_path, embeddings)
    logger.info(f"Database FAISS caricato da {db_path}")
except Exception as e:
    logger.error(f"Errore nel caricamento del database FAISS: {e}")
    faiss_index = None

# Inizializzazione FastAPI
app = FastAPI()

# Configurazione CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# *********************************************************************
# Funzione per gestire il modello LLM dinamicamente

def get_llm():
    if ACTIVE_MODEL == "gpt-4o-mini":
        return ChatOpenAI(model="gpt-4o-mini", openai_api_key=OPENAI_API_KEY)
    # Decomentare per utilizzare i modelli Anthropic:
    # elif ACTIVE_MODEL == "claude-v1":
    #     return ChatAnthropic(model="claude-v1", anthropic_api_key=ANTHROPIC_API_KEY)
    # elif ACTIVE_MODEL == "claude-3.5-haiku":
    #     return ChatAnthropic(model="claude-3.5-haiku", anthropic_api_key=ANTHROPIC_API_KEY)
    else:
        raise ValueError("Modello non supportato.")

# *********************************************************************
# Funzione per salvare file caricati

def save_file(file: UploadFile, folder: str) -> str:
    try:
        file_path = os.path.join(folder, file.filename)
        with open(file_path, "wb") as f:
            f.write(file.file.read())
        return file_path
    except Exception as e:
        logger.error(f"Errore durante il salvataggio del file {file.filename}: {e}")
        raise HTTPException(status_code=500, detail=f"Errore nel salvataggio del file {file.filename}")

# *********************************************************************
# Funzione per estrarre contenuto dai file e dividerlo in chunk

def extract_content(file_path: str) -> List[Document]:
    try:
        if file_path.endswith(".pdf"):
            reader = PdfReader(file_path)
            text = [page.extract_text() for page in reader.pages if page.extract_text()]
        elif file_path.endswith(".docx"):
            doc = DocxDocument(file_path)
            text = [para.text for para in doc.paragraphs if para.text.strip()]
        elif file_path.endswith(".pptx"):
            ppt = Presentation(file_path)
            text = []
            for slide in ppt.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
        elif file_path.endswith(".txt"):
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.readlines()
        else:
            raise ValueError("Formato file non supportato")

        # Combina il contenuto in un unico testo e suddividilo in chunk
        text_combined = "\n".join(text)
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=100)
        chunks = text_splitter.split_text(text_combined)

        return [Document(page_content=chunk) for chunk in chunks]

    except Exception as e:
        logger.error(f"Errore nell'estrazione del contenuto da {file_path}: {e}")
        raise HTTPException(status_code=500, detail=f"Errore nell'estrazione del contenuto da {file_path}")

# *********************************************************************
# Funzione per recuperare contesto dal database FAISS

def retrieve_context(query: str) -> List[dict]:
    if not faiss_index:
        logger.warning("Database FAISS non caricato. Nessun contesto disponibile.")
        return []
    try:
        docs = faiss_index.similarity_search(query, k=5)
        logger.info(f"Documenti recuperati: {docs}")
        return [
            {
                "content": doc.page_content,
                "source": {
                    "filename": doc.metadata.get("filename", "sconosciuto"),
                    "page_number": doc.metadata.get("page_number", "n/a")
                }
            }
            for doc in docs
        ]
    except Exception as e:
        logger.error(f"Errore nel recupero del contesto: {str(e)}")
        return []

# *********************************************************************
# Endpoint per il completamento della chat

class ChatRequest(BaseModel):
    messages: List[dict]
    temperature: float = 0.7

@app.post("/v1/chat/completions")
async def chat_completion(request: ChatRequest):
    try:
        llm = get_llm()
        response = llm.predict_messages(messages=request.messages, temperature=request.temperature)
        return {"llm_response": response.content}
    except Exception as e:
        logger.error(f"Errore nel completamento della chat: {e}")
        raise HTTPException(status_code=500, detail=f"Errore nel completamento della chat: {e}")

# *********************************************************************
# Avvio del server

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)